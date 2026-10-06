import json
import os
from openpyxl import load_workbook

EXCEL_PATH = "data/tu_vung_hsk.xlsx"
OUTPUT_PATH = "data/vocab_meanings.json"

if not os.path.exists(EXCEL_PATH):
    print(f"ERROR: Khong tim thay {EXCEL_PATH}")
    exit(1)

print(f"Dang doc: {EXCEL_PATH}")
wb = load_workbook(EXCEL_PATH, data_only=True, read_only=True)
result = {}

for sheet_name in wb.sheetnames:
    if not sheet_name.upper().startswith("HSK"):
        continue
    ws = wb[sheet_name]
    is_hsk79 = any(x in sheet_name for x in ['7', '8', '9'])
    col_zh = 1
    col_vi = 4 if is_hsk79 else 5
    count_sheet = 0
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or len(row) <= col_vi:
            continue
        zh = str(row[col_zh]).strip() if row[col_zh] else ""
        vi = str(row[col_vi]).strip() if len(row) > col_vi and row[col_vi] else ""
        if zh and vi and zh not in result:
            result[zh] = vi
            count_sheet += 1
    print(f"  {sheet_name}: {count_sheet} tu")

wb.close()

with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, separators=(",", ":"))

size_kb = os.path.getsize(OUTPUT_PATH) / 1024
print(f"OK: {len(result)} tu, {size_kb:.1f} KB")
print(f"Output: {OUTPUT_PATH}")
